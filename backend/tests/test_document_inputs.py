"""Tests for PDF/DOCX/email input extraction and upload validation."""

from __future__ import annotations

import io
import tempfile
import unittest
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from fastapi.testclient import TestClient

from ul.document_chunking import extract_pages_from_pdf
from ul.document_extractor import (
    document_type_for_filename,
    extract_pages_from_document,
    extract_pages_from_docx,
    extract_pages_from_email,
    is_supported_upload_filename,
)
from ul.ul_service import (
    _resolve_source_paths,
    get_reference_document_paths,
    process_documents,
    save_uploaded_files,
)


def _write_pdf(path: Path, text: str) -> None:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    doc.save(path)
    doc.close()


def _write_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Product Technical Specification", level=1)
    document.add_paragraph("Model iPhone 16 with USB-C connector.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Part number"
    table.cell(1, 0).text = "Battery"
    table.cell(1, 1).text = "A2991"
    document.save(path)


def _write_eml(path: Path, *, attach_pdf: bool = False) -> None:
    message = EmailMessage()
    message["Subject"] = "Certification request"
    message["From"] = "oem@example.com"
    message["To"] = "ul@example.com"
    message["Cc"] = "pm@example.com"
    message.set_content(
        "Product = iPhone 16\nCertification request for US launch."
    )
    if attach_pdf:
        pdf_bytes = _pdf_bytes("Battery datasheet rating 3.8V")
        message.add_attachment(
            pdf_bytes,
            maintype="application",
            subtype="pdf",
            filename="battery_datasheet.pdf",
        )
    path.write_bytes(bytes(message))


def _pdf_bytes(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    payload = doc.tobytes()
    doc.close()
    return payload


class UploadValidationTests(unittest.TestCase):
    def test_accepts_pdf_docx_and_email_extensions(self) -> None:
        self.assertTrue(is_supported_upload_filename("spec.pdf"))
        self.assertTrue(is_supported_upload_filename("BOM.DOCX"))
        self.assertTrue(is_supported_upload_filename("request.eml"))
        self.assertTrue(is_supported_upload_filename("request.msg"))

    def test_rejects_unsupported_extensions(self) -> None:
        self.assertFalse(is_supported_upload_filename("notes.txt"))
        self.assertFalse(is_supported_upload_filename("sheet.xlsx"))
        self.assertFalse(is_supported_upload_filename(None))
        with self.assertRaises(ValueError):
            document_type_for_filename("notes.txt")

    def test_document_types(self) -> None:
        self.assertEqual(document_type_for_filename("a.pdf"), "pdf")
        self.assertEqual(document_type_for_filename("a.docx"), "docx")
        self.assertEqual(document_type_for_filename("a.eml"), "email")
        self.assertEqual(document_type_for_filename("a.msg"), "email")

    def test_save_uploaded_files_accepts_supported_types(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            with patch("ul.ul_service.config.UPLOAD_DIR", tmp_path):
                result = save_uploaded_files(
                    [
                        ("spec.pdf", b"%PDF-1.4"),
                        ("bom.docx", b"PK"),
                        ("request.eml", b"From: a@b.c\n\nhello"),
                    ]
                )
            saved = set(result["files"])
            self.assertEqual(saved, {"spec.pdf", "bom.docx", "request.eml"})
            self.assertTrue((tmp_path / result["document_id"] / "spec.pdf").exists())

    def test_save_uploaded_files_rejects_unsupported_type(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            with patch("ul.ul_service.config.UPLOAD_DIR", Path(tmp)):
                with self.assertRaises(ValueError):
                    save_uploaded_files([("notes.txt", b"hello")])

    def test_resolve_source_paths_finds_all_supported_types(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            upload_dir = Path(tmp) / "doc-1"
            upload_dir.mkdir()
            (upload_dir / "a.pdf").write_bytes(b"x")
            (upload_dir / "b.docx").write_bytes(b"x")
            (upload_dir / "c.eml").write_bytes(b"x")
            (upload_dir / "ignore.txt").write_bytes(b"x")
            with patch("ul.ul_service.config.UPLOAD_DIR", Path(tmp)):
                paths = _resolve_source_paths("doc-1", None)
            names = [path.name for path in paths]
            self.assertEqual(names, ["a.pdf", "b.docx", "c.eml"])

    def test_upload_api_rejects_unsupported_extension(self) -> None:
        from main import app

        client = TestClient(app)
        response = client.post(
            "/api/ul/upload",
            files=[("files", ("notes.txt", io.BytesIO(b"hello"), "text/plain"))],
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("PDF, DOCX, or email", response.json()["detail"])

    def test_upload_api_accepts_docx_and_email(self) -> None:
        from main import app

        client = TestClient(app)
        with tempfile.TemporaryDirectory() as tmp:
            with patch("ul.ul_service.config.UPLOAD_DIR", Path(tmp)):
                response = client.post(
                    "/api/ul/upload",
                    files=[
                        ("files", ("spec.docx", io.BytesIO(b"PK"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                        ("files", ("request.eml", io.BytesIO(b"From: a@b.c\n\nbody"), "message/rfc822")),
                    ],
                )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertIn("document_id", body)
        self.assertEqual(set(body["files"]), {"spec.docx", "request.eml"})


class PdfExtractionCompatibilityTests(unittest.TestCase):
    def test_pdf_extractor_still_returns_source_file_page_and_text(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "product.pdf"
            _write_pdf(path, "iPhone 16 electrical characteristics")
            pages = extract_pages_from_pdf(str(path))

        self.assertGreaterEqual(len(pages), 1)
        page = pages[0]
        self.assertEqual(page["source_file"], "product.pdf")
        self.assertEqual(page["page_number"], 1)
        self.assertIn("iPhone 16", page["text"])
        self.assertNotIn("document_type", page)

    def test_dispatcher_adds_document_type_without_changing_pdf_text(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "product.pdf"
            _write_pdf(path, "iPhone 16 electrical characteristics")
            raw = extract_pages_from_pdf(str(path))
            wrapped = extract_pages_from_document(str(path))

        self.assertEqual(wrapped[0]["text"], raw[0]["text"])
        self.assertEqual(wrapped[0]["page_number"], raw[0]["page_number"])
        self.assertEqual(wrapped[0]["source_file"], raw[0]["source_file"])
        self.assertEqual(wrapped[0]["document_type"], "pdf")


class DocxExtractionTests(unittest.TestCase):
    def test_extracts_headings_paragraphs_and_tables(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "spec.docx"
            _write_docx(path)
            pages = extract_pages_from_docx(str(path))

        self.assertEqual(len(pages), 1)
        page = pages[0]
        self.assertEqual(page["source_file"], "spec.docx")
        self.assertEqual(page["page_number"], 1)
        self.assertEqual(page["document_type"], "docx")
        self.assertIn("# Product Technical Specification", page["text"])
        self.assertIn("Model iPhone 16", page["text"])
        self.assertIn("[Table]", page["text"])
        self.assertIn("Battery | A2991", page["text"])

    def test_dispatcher_routes_docx(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "spec.docx"
            _write_docx(path)
            pages = extract_pages_from_document(str(path))
        self.assertEqual(pages[0]["document_type"], "docx")


class EmailExtractionTests(unittest.TestCase):
    def test_extracts_headers_and_body(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "request.eml"
            _write_eml(path)
            pages = extract_pages_from_email(str(path))

        self.assertGreaterEqual(len(pages), 1)
        text = pages[0]["text"]
        self.assertEqual(pages[0]["source_file"], "request.eml")
        self.assertEqual(pages[0]["page_number"], 1)
        self.assertEqual(pages[0]["document_type"], "email")
        self.assertIn("Subject: Certification request", text)
        self.assertIn("From: oem@example.com", text)
        self.assertIn("To: ul@example.com", text)
        self.assertIn("Cc: pm@example.com", text)
        self.assertIn("Product = iPhone 16", text)

    def test_includes_pdf_attachment_content_as_later_pages(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "request.eml"
            _write_eml(path, attach_pdf=True)
            pages = extract_pages_from_email(str(path))

        self.assertGreaterEqual(len(pages), 2)
        self.assertIn("Attachments:", pages[0]["text"])
        self.assertIn("battery_datasheet.pdf", pages[0]["text"])
        attachment = pages[1]
        self.assertEqual(attachment["source_file"], "request.eml")
        self.assertEqual(attachment["page_number"], 2)
        self.assertEqual(attachment["document_type"], "email")
        self.assertIn("[Attachment: battery_datasheet.pdf]", attachment["text"])
        self.assertIn("Battery datasheet", attachment["text"])


class ProcessDocumentsIndexingTests(unittest.TestCase):
    def test_indexes_all_chunks_before_ul_extraction(self) -> None:
        chunks = [
            {
                "chunk_id": "chunk_000001",
                "source_file": "product.pdf",
                "page_range": {"start": 1, "end": 1},
                "text": "Product text",
            }
        ]
        graph_data = {
            "companies": [],
            "products": [],
            "models": [],
            "components": [],
            "normalized_components": [],
            "standards": [],
            "clauses": [],
            "certifications": [],
            "tests": [],
            "relationships": [],
        }
        call_order: list[str] = []

        def _index(all_chunks, run_id=None):
            call_order.append("index")
            self.assertEqual(all_chunks, chunks)
            self.assertEqual(run_id, "doc-1")
            self.assertEqual(all_chunks[0]["chunk_id"], "chunk_000001")
            return 1

        def _extract(all_chunks, target_product_name=None, project_id=None):
            call_order.append("extract")
            self.assertEqual(all_chunks, chunks)
            self.assertEqual(project_id, "doc-1")
            self.assertEqual(all_chunks[0]["project_id"], "doc-1")
            self.assertEqual(all_chunks[0]["metadata"]["project_id"], "doc-1")
            return graph_data

        with (
            patch("ul.ul_service._resolve_source_paths", return_value=[Path("product.pdf")]),
            patch(
                "ul.ul_service._process_single_document",
                return_value=(chunks, 1, 2, 1),
            ),
            patch("ul.ul_service.get_reference_document_paths", return_value=[]),
            patch("ul.ul_service.index_chunks_with_lightrag", side_effect=_index),
            patch("ul.ul_service.extract_graph_from_chunks", side_effect=_extract),
            patch("ul.ul_service.write_triplets_csv", return_value="triplets.csv"),
            patch("ul.ul_service.ingest_triplets"),
        ):
            result = process_documents(document_id="doc-1")

        self.assertEqual(call_order, ["index", "extract"])
        self.assertEqual(result["extraction_engine"], "lightrag.prompt")
        self.assertEqual(result["vector_storage"], "LightRAG")
        self.assertEqual(result["warnings"], [])

    def test_extracts_from_uploads_and_reference_documents(self) -> None:
        user_chunks = [
            {
                "chunk_id": "chunk_000001",
                "source_file": "product.pdf",
                "page_range": {"start": 1, "end": 1},
                "text": "Product is evaluated against IEC 62368-1:2023.",
            }
        ]
        reference_chunks = [
            {
                "chunk_id": "ref_standards_000001",
                "source_file": "standards.pdf",
                "page_range": {"start": 1, "end": 1},
                "text": "TEST-001 CERT-001 Clause 6 Annex G",
            }
        ]
        graph_data = {
            "companies": [],
            "products": [],
            "models": [],
            "components": [],
            "normalized_components": [],
            "standards": [],
            "clauses": [],
            "certifications": [],
            "tests": [],
            "relationships": [],
        }
        indexed: list[dict] = []
        extracted: list[dict] = []

        def _process(path, *, chunk_start_index, chunk_id_prefix="chunk"):
            del chunk_start_index, chunk_id_prefix
            if Path(path).name == "product.pdf":
                return (user_chunks, 1, 2, 1)
            return (reference_chunks, 1, 2, 1)

        def _index(chunks, run_id=None):
            indexed.extend(chunks)
            return len(chunks)

        def _extract(chunks, target_product_name=None, project_id=None):
            extracted.extend(chunks)
            self.assertEqual(project_id, "doc-1")
            return graph_data

        with (
            patch("ul.ul_service._resolve_source_paths", return_value=[Path("product.pdf")]),
            patch("ul.ul_service._process_single_document", side_effect=_process),
            patch(
                "ul.ul_service.get_reference_document_paths",
                return_value=[Path("standards.pdf")],
            ),
            patch("ul.ul_service.index_chunks_with_lightrag", side_effect=_index),
            patch("ul.ul_service.extract_graph_from_chunks", side_effect=_extract),
            patch("ul.ul_service.write_triplets_csv", return_value="triplets.csv"),
            patch("ul.ul_service.ingest_triplets"),
        ):
            process_documents(document_id="doc-1")

        self.assertEqual(len(indexed), 2)
        self.assertEqual(
            {c["chunk_id"] for c in indexed},
            {"chunk_000001", "ref_standards_000001"},
        )
        self.assertEqual(
            [c["chunk_id"] for c in extracted],
            ["chunk_000001", "ref_standards_000001"],
        )
        self.assertIn("TEST-001", extracted[1]["text"])
        self.assertEqual(extracted[0]["project_id"], "doc-1")
        self.assertEqual(extracted[1]["project_id"], "doc-1")

    def test_reference_folder_lists_all_supported_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            (folder / "README.md").write_text("ignore", encoding="utf-8")
            (folder / "notes.txt").write_text("ignore", encoding="utf-8")
            (folder / "extra.pdf").write_bytes(b"%PDF-1.4")
            (folder / "standards.pdf").write_bytes(b"%PDF-1.4")
            with patch("ul.ul_service.config.REFERENCE_DOCUMENTS_DIR", folder):
                names = [path.name for path in get_reference_document_paths()]
        self.assertEqual(names[0], "standards.pdf")
        self.assertIn("extra.pdf", names)
        self.assertNotIn("README.md", names)
        self.assertNotIn("notes.txt", names)


if __name__ == "__main__":
    unittest.main()
